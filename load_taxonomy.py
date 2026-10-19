#!/usr/bin/env python3
"""
Load Chadia's authoritative complaint taxonomy and re-flag "new codes" against it.

WHY: the classifier's output was checked against the legacy `cx_complaints` table,
which is NOT the taxonomy the business uses. Chadia supplied the real one:
  - Complaint Validation Audit - Generic Set of Complaints .csv     (19 codes)
  - Complaint Validation Audit - Non generic Complaints Types.csv   (115 codes)
  - Complaint Validation Audit .xlsx -> 'Audit Suggestions'         (approved renames)
  - data/vocab_additions.csv                                        (4 codes she approved Sep 22)
Result of the wrong reference: 15 of 27 codes flagged "new" were codes she had
already defined (857 reviews), and 9 more were variants of a code she approved
renaming to "Not as Advertised/Misleading".

This builds `complaint_codes` from her files, then recomputes is_new_code and the
proposed_codes queue against it.

DECISIONS ARE DURABLE (2026-10-18). Approve/deny used to live only in
proposed_codes.status while this script DELETEd and rebuilt that table every 30
minutes, so her ruling on a new category reverted within one refresh. Now:
  * complaint_codes rows with source 'chadia-approved' / 'chadia-denied' are her
    decision and are never overwritten or deleted here;
  * an approved code is part of the known vocabulary (its rows stop being "new");
  * a denied code is excluded from the vocabulary AND from the queue, and
    sync_corrections.py remaps the affected classifications to the code she meant;
  * proposed_codes is upserted, never truncated, so status/decided_at survive.

Idempotent: safe to run on every refresh.
"""
import csv, os, re
import openpyxl, pymysql

ATT = "/opt/data/attachments"
GENERIC = f"{ATT}/Complaint Validation Audit  - Generic Set of Complaints .csv"
NONGEN = f"{ATT}/Complaint Validation Audit  - Non generic Complaints Types.csv"
AUDIT = f"{ATT}/Complaint Validation Audit .xlsx"
ADDITIONS = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                         "data", "vocab_additions.csv")

DECISION_SOURCES = ("chadia-approved", "chadia-denied")

# Her approved rename target -> the label the classifier emits, so a variant is not
# reported as a brand-new code. Kept explicit: string-normalising alone misfires.
ALIASES = {
    "not as described": "Not as Advertised/Misleading",
    "not as described online cannot use": "Not as Advertised/Misleading",
    "not as described as per functionality": "Not as Advertised/Misleading",
    "not as described on website": "Not as Advertised/Misleading",
    "received wrong item": "Not as Advertised/Misleading",
    "not as expected": "Not as Advertised/Misleading",
}

# Two rename targets in her Audit Suggestions sheet are spellings that do not
# exist as codes, and inserting them creates duplicates of codes she already owns
# ('Noisey' next to 'Noisy', 'Sharp and/or Cut Finger' next to 'Sharp'). They are
# aliases, not new codes.
PHANTOM_TARGETS = {
    "noisey": "Noisy",
    "sharp and/or cut finger": "Sharp",
    "sharp and/or cut finger ": "Sharp",
    "sharp and or cut finger": "Sharp",
}

# Codes her file leaves with no category and the legacy fallback cannot resolve.
# Mirrors cxvocab.CATEGORY_OVERRIDE; kept here too because this script is the one
# that writes complaint_codes.category (and therefore the 519 NULL-category rows).
CATEGORY_OVERRIDE = {
    "too short": "Size",
    "too tall": "Size",
    "too loose": "Size",
    "strap too long": "Size",
    "difficult to use": "Function",
    "not adjustable": "Function",
    "sticky": "Material",
    "pilling": "Material",
    "velcro damage": "Material",
    "poor insulation": "Quality",
}


def norm(s):
    return re.sub(r"\s+", " ", (s or "").strip().lower())


def load_her_codes():
    rows = {}
    for path, src in ((GENERIC, "generic"), (NONGEN, "non-generic"),
                      (ADDITIONS, "approved-addition")):
        if not os.path.exists(path):
            print(f"WARN: missing source {path}")
            continue
        for i, r in enumerate(csv.reader(open(path, encoding="utf-8-sig"))):
            if i == 0 or not r or not (r[0] or "").strip():
                continue
            code = r[0].strip()
            cat = (r[1] or "").strip() or CATEGORY_OVERRIDE.get(norm(code))
            rows[norm(code)] = {"code": code, "category": cat, "source": src, "approved": 1}
    return rows


def read_renames(rows, aliases):
    """Fold her approved renames into the code set and the alias table."""
    wb = openpyxl.load_workbook(AUDIT, data_only=True)
    ws = wb["Audit Suggestions"]
    n = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or len(r) < 3 or not r[2]:
            continue
        m = re.search(r"change to ['\"](.+?)['\"]", str(r[2]), re.I)
        if not m:
            continue
        n += 1
        target = m.group(1).strip()
        tn = norm(target)
        if tn in PHANTOM_TARGETS:                    # duplicate spelling -> real code
            target = PHANTOM_TARGETS[tn]
            tn = norm(target)
        src_cat = (r[1] or "").strip() or None
        aliases[norm(r[0])] = target
        if tn not in rows:
            # The target's OWN category where her file has one; otherwise the
            # source row's. Inheriting blindly is how 'Noisey' became Transportation.
            rows[tn] = {"code": target,
                        "category": CATEGORY_OVERRIDE.get(tn) or src_cat,
                        "source": "rename", "approved": 1}
    return n


def main():
    rows = load_her_codes()
    aliases = dict(ALIASES)
    n_renames = read_renames(rows, aliases)
    print(f"her taxonomy: {len(rows)} distinct codes, {n_renames} rename rules")

    conn = pymysql.connect(host=os.environ["CXAPP_DB_HOST"], user=os.environ["CXAPP_DB_USER"],
                           password=os.environ["CXAPP_DB_PASSWORD"], database=os.environ["CXAPP_DB_NAME"],
                           charset="utf8mb4", autocommit=True)
    cur = conn.cursor()
    cur.execute("""CREATE TABLE IF NOT EXISTS complaint_codes (
                     code VARCHAR(120) NOT NULL,
                     category VARCHAR(80) NULL,
                     source VARCHAR(20) NOT NULL,
                     approved TINYINT(1) NOT NULL DEFAULT 1,
                     PRIMARY KEY (code)) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""")

    # --- her existing decisions -------------------------------------------------
    cur.execute("SELECT code, category, source, approved FROM complaint_codes")
    decisions = {norm(c): {"code": c, "category": cat, "source": s, "approved": int(a)}
                 for c, cat, s, a in cur.fetchall() if s in DECISION_SOURCES}
    approved_extra = {k for k, d in decisions.items() if d["approved"]}
    denied = {k for k, d in decisions.items() if not d["approved"]}

    # stale rows from her files (e.g. the old 'Noisey' phantom) are removed; her
    # decisions never are.
    keep = [v["code"] for v in rows.values()] + [d["code"] for d in decisions.values()]
    placeholders = ",".join(["%s"] * len(keep))
    cur.execute(f"""DELETE FROM complaint_codes
                    WHERE source IN ('generic','non-generic','rename','approved-addition')
                      AND code NOT IN ({placeholders})""", keep)

    for k, v in rows.items():
        if k in decisions:                          # her ruling outranks the file
            continue
        cur.execute("""INSERT INTO complaint_codes (code,category,source,approved)
                       VALUES (%s,%s,%s,%s)
                       ON DUPLICATE KEY UPDATE category=VALUES(category),
                                               source=VALUES(source), approved=VALUES(approved)""",
                    (v["code"], v["category"], v["source"], v["approved"]))
    cur.execute("SELECT COUNT(*), SUM(approved=0) FROM complaint_codes")
    tot, den = cur.fetchone()
    print(f"complaint_codes rows: {tot} (approved={tot - int(den or 0)}, denied={int(den or 0)})")

    # Her file leaves 91 codes with no category. Where the classified rows already
    # know one, adopt the modal value: without this, an approved/denied code gets
    # remapped without its category and the row lands in the wrong rollup.
    cur.execute("""UPDATE complaint_codes cc
                   JOIN (SELECT complaint,
                                SUBSTRING_INDEX(GROUP_CONCAT(category ORDER BY n DESC), ',', 1) AS category
                           FROM (SELECT complaint, category, COUNT(*) n
                                   FROM classifications
                                  WHERE category IS NOT NULL AND category NOT IN ('','NEW','ERR','NO_TEXT')
                                  GROUP BY complaint, category) t
                          GROUP BY complaint) x ON x.complaint = cc.code
                      SET cc.category = x.category
                    WHERE cc.category IS NULL""")
    print(f"code categories adopted from classified rows: {cur.rowcount}")

    known = (set(rows) | approved_extra) - denied

    # --- re-flag ---------------------------------------------------------------
    # An alias must RESOLVE to a code she owns, not merely be excluded: otherwise
    # "Received Wrong Item" still shows as a brand-new code and asks her to decide
    # something she has already decided.
    def resolve(c):
        n = norm(c)
        return norm(aliases.get(n, n))

    cur.execute("SELECT review_id, complaint FROM classifications")
    all_rows = cur.fetchall()
    is_new, still_denied = [], []
    for rid, c in all_rows:
        if not c:
            continue
        n = norm(c)
        if n in denied:
            still_denied.append(rid)                # should be 0: deny remaps them
        elif resolve(c) not in known:
            is_new.append(rid)

    cur.execute("UPDATE classifications SET is_new_code=0")
    if is_new:
        for i in range(0, len(is_new), 1000):
            chunk = is_new[i:i + 1000]
            cur.execute(f"UPDATE classifications SET is_new_code=1 WHERE review_id IN ({','.join(map(str,chunk))})")
    print(f"rows using a code outside her taxonomy: {len(is_new)}")
    if still_denied:
        print(f"WARN: {len(still_denied)} rows still carry a DENIED code (deny did not remap them)")

    # --- categories: fill blanks from the taxonomy, never overwrite a real one ---
    cur.execute("""UPDATE classifications c JOIN complaint_codes cc ON cc.code = c.complaint
                      SET c.category = cc.category
                    WHERE (c.category IS NULL OR c.category='') AND cc.category IS NOT NULL""")
    print(f"categories backfilled: {cur.rowcount}")
    cur.execute("""SELECT COUNT(*) FROM classifications WHERE category IS NULL OR category=''""")
    print(f"rows still without a category: {cur.fetchone()[0]}")

    # --- proposed_codes: upsert, never truncate --------------------------------
    cur.execute("""SELECT complaint, COUNT(*) n FROM classifications
                   WHERE is_new_code=1 AND complaint IS NOT NULL
                   GROUP BY complaint ORDER BY n DESC""")
    real = [(c, n) for c, n in cur.fetchall() if norm(c) not in denied]

    for code, n in real:
        cur.execute("""INSERT INTO proposed_codes (code, n_rows, example_review, status)
                       VALUES (%s, %s,
                               (SELECT LEFT(review_text,400) FROM classifications WHERE complaint=%s LIMIT 1),
                               'pending')
                       ON DUPLICATE KEY UPDATE
                         n_rows=VALUES(n_rows),
                         example_review=IF(status='pending', VALUES(example_review), example_review)""",
                    (code, n, code))

    if real:
        ph = ",".join(["%s"] * len(real))
        cur.execute(f"DELETE FROM proposed_codes WHERE status='pending' AND code NOT IN ({ph})",
                    [c for c, _ in real])
    else:
        cur.execute("DELETE FROM proposed_codes WHERE status='pending'")
    print(f"genuinely new codes remaining: {len(real)}")
    for code, n in real:
        print(f"   {n:5d}  {code}")

    cur.execute("SELECT status, COUNT(*) FROM proposed_codes GROUP BY status")
    print("proposed_codes:", dict(cur.fetchall()))
    conn.close()


if __name__ == "__main__":
    main()
